"""Extraction locale des pièces, sans exécution de leur contenu."""
import base64
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

from runtime_paths import ROOT, worker_command, hidden_process
RUNTIME = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies'
MAX_FILE = 12 * 1024 * 1024
MAX_TOTAL = 30 * 1024 * 1024
MAX_TEXT = 140000
TEXT_TYPES = {'.txt', '.md', '.csv', '.tsv', '.json', '.xml', '.html', '.log', '.py', '.js', '.ts', '.css', '.sql', '.yaml', '.yml', '.svg', '.ics', '.tex', '.sh'}
TYPES = TEXT_TYPES | {'.pdf', '.docx', '.xlsx', '.pptx', '.png', '.jpg', '.jpeg', '.zip'}


def safe_name(name):
    name = str(name).replace('\\', '/').split('/')[-1]
    return ''.join(c for c in name if c.isprintable() and c not in '\r\n')[:160] or 'piece.txt'


def decode_upload(data):
    name = safe_name(data.get('name', ''))
    if Path(name).suffix.lower() not in TYPES:
        raise ValueError('Format non pris en charge : ' + name + '. Utilise PDF, Word .docx, Excel .xlsx, PowerPoint .pptx, texte, PNG ou JPEG.')
    encoded = data.get('content', '')
    if not isinstance(encoded, str) or len(encoded) > MAX_FILE * 4 // 3 + 8:
        raise ValueError('Chaque pièce doit peser au maximum 12 Mo.')
    try:
        raw = base64.b64decode(encoded, validate=True)
    except (ValueError, TypeError):
        raise ValueError('Le fichier n’a pas pu être reçu. Sélectionne-le à nouveau.')
    if not raw or len(raw) > MAX_FILE:
        raise ValueError('La pièce est vide ou dépasse 12 Mo.')
    return name, raw


def extract(name, raw):
    """Le parseur travaille dans un processus borné ; seuls les résultats reviennent."""
    with tempfile.TemporaryDirectory(prefix='duo-piece-') as folder:
        path = Path(folder) / ('piece' + Path(name).suffix.lower())
        path.write_bytes(raw)
        try:
            output = Path(folder) / 'result.json'
            result = subprocess.run(worker_command('extract', path, output), capture_output=True, text=True, timeout=120, **hidden_process())
            data = json.loads(output.read_text(encoding='utf-8'))
        except (OSError, subprocess.TimeoutExpired, ValueError):
            raise ValueError('Lecture impossible ou trop longue pour ' + name + '. Essaie une version plus légère.')
        if result.returncode or data.get('error'):
            raise ValueError(name + ' : ' + data.get('error', 'lecture impossible'))
        text = data['text'].strip()
        if not text:
            raise ValueError(name + ' : aucun texte lisible détecté.')
        if len(text) > MAX_TEXT:
            raise ValueError(name + ' dépasse la capacité de lecture de ce dossier. Divise le document ; aucun extrait tronqué n’a été envoyé.')
        return {'text': text, 'warnings': data.get('warnings', []), 'sha256': hashlib.sha256(raw).hexdigest()}


def ocr(path):
    if os.name == 'nt':
        result = subprocess.run(['powershell.exe', '-NoProfile', '-NonInteractive', '-File', str(ROOT / 'desktop/ocr-windows.ps1'), str(Path(path).resolve())], capture_output=True, timeout=40, **hidden_process())
        if result.returncode:
            raise ValueError('La reconnaissance de texte de Windows n’a pas pu lire cette image. Vérifie les langues installées ou fournis un document avec du texte sélectionnable.')
        return result.stdout.decode('utf-8-sig').strip()
    binary = ROOT / 'bin/duo-ocr'
    if not binary.is_file():
        raise ValueError('La lecture des images n’est pas installée. Utilise un document contenant du texte sélectionnable.')
    result = subprocess.run([str(binary), str(path)], capture_output=True, text=True, timeout=40, **hidden_process())
    if result.returncode:
        raise ValueError('La reconnaissance du texte de cette image a échoué.')
    return result.stdout.strip()


def parse(path):
    suffix = path.suffix.lower()
    warnings = []
    if suffix == '.zip':
        sections = []
        with zipfile.ZipFile(path) as archive:
            entries = [e for e in archive.infolist() if not e.is_dir() and not e.filename.startswith('__MACOSX/') and not e.filename.endswith('.DS_Store')]
            if not 1 <= len(entries) <= 30 or sum(e.file_size for e in entries) > MAX_TOTAL:
                raise ValueError('L’archive dépasse 30 fichiers ou 30 Mo après décompression.')
            for index, entry in enumerate(entries):
                extension = Path(entry.filename).suffix.lower()
                if extension not in TYPES or extension == '.zip' or entry.file_size > MAX_FILE:
                    raise ValueError('Le fichier interne ' + safe_name(entry.filename) + ' est trop volumineux ou non pris en charge. Les ZIP imbriqués ne sont pas acceptés.')
                with tempfile.TemporaryDirectory(prefix='duo-zip-piece-') as folder:
                    child = Path(folder) / ('piece' + extension)
                    child.write_bytes(archive.read(entry))
                    parsed = parse(child)
                    sections.append('[Fichier de l’archive : ' + entry.filename + ']\n' + parsed['text'])
                    warnings += [safe_name(entry.filename) + ' : ' + w for w in parsed.get('warnings', [])]
                if sum(map(len, sections)) > MAX_TEXT:
                    raise ValueError('Le texte de l’archive dépasse la capacité du dossier. Aucun contenu tronqué n’a été envoyé.')
        return {'text': '\n\n'.join(sections), 'warnings': warnings}
    if suffix in TEXT_TYPES:
        raw = path.read_bytes()
        for encoding in ('utf-8-sig', 'utf-16', 'cp1252'):
            try:
                text = raw.decode(encoding)
                if '\x00' in text:
                    continue
                return {'text': text, 'warnings': warnings}
            except UnicodeError:
                continue
        raise ValueError('Encodage de texte non reconnu.')
    if suffix in ('.png', '.jpg', '.jpeg'):
        return {'text': ocr(path), 'warnings': ['Reconnaissance de texte sur image : des erreurs sont possibles. L’aspect visuel et les graphiques ne sont pas évalués.']}
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('PDF protégé par mot de passe. Fournis une copie déverrouillée.')
        if len(reader.pages) > 100:
            raise ValueError('Ce PDF dépasse 100 pages. Divise-le pour garantir une lecture complète.')
        pages = []
        for index, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            if not text.strip():
                with tempfile.TemporaryDirectory(prefix='duo-page-') as folder:
                    prefix = Path(folder) / 'scan'
                    import pypdfium2 as pdfium
                    pdf = pdfium.PdfDocument(str(path))
                    try:
                        rendered = pdf[index - 1]
                        bitmap = rendered.render(scale=min(3, 2000 / max(rendered.get_size())))
                        bitmap.to_pil().save(prefix.with_suffix('.png'))
                        bitmap.close();rendered.close()
                    finally:
                        pdf.close()
                    text = ocr(prefix.with_suffix('.png'))
                    if not text:
                        text = '[Page sans texte détectable : contenu visuel non analysé.]'
                    warnings.append('Page ' + str(index) + ' : texte reconnu sur image, à vérifier.')
            pages.append('[Page ' + str(index) + ']\n' + text)
            if sum(map(len, pages)) > MAX_TEXT:
                raise ValueError('Le texte dépasse la capacité de lecture. Divise le PDF ; aucune page n’a été envoyée.')
        warnings.append('Lecture du texte : la mise en page, les images, signatures et graphiques ne sont pas vérifiés.')
        return {'text': '\n\n'.join(pages), 'warnings': warnings}
    # Contrôle de taille avant de décompresser les formats Office.
    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 5000 or sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
            raise ValueError('Document trop volumineux après décompression.')
        if suffix == '.docx':
            sections = []
            names = ['word/document.xml'] + sorted(n for n in archive.namelist() if n.startswith(('word/header', 'word/footer', 'word/footnotes', 'word/endnotes')) and n.endswith('.xml'))
            for name in names:
                root = ET.fromstring(archive.read(name))
                ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
                lines = [''.join(node.text or '' for node in p.iter(ns + 't')) for p in root.iter(ns + 'p')]
                sections.append('[' + name.split('/')[-1] + ']\n' + '\n'.join(line for line in lines if line))
            warnings.append('Lecture du texte, y compris les cellules des tableaux. Images, commentaires et présentation non évalués.')
            return {'text': '\n\n'.join(sections), 'warnings': warnings}
        if suffix == '.pptx':
            names = sorted((n for n in archive.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml') and '/_rels/' not in n), key=lambda n: int(n.rsplit('slide', 1)[1].split('.')[0]))
            sections = []
            for index, name in enumerate(names, 1):
                root = ET.fromstring(archive.read(name))
                text = '\n'.join(n.text or '' for n in root.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'))
                sections.append('[Diapositive ' + str(index) + ']\n' + text)
            return {'text': '\n\n'.join(sections), 'warnings': ['Texte des diapositives uniquement : visuels, graphiques et notes ne sont pas évalués.']}
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        lines = []
        length = 0
        try:
            for sheet in workbook:
                lines.append('[Feuille : ' + sheet.title + ']')
                if sheet.max_row > 20000 or sheet.max_column > 200:
                    raise ValueError('Une feuille dépasse 20 000 lignes ou 200 colonnes. Exporte la zone utile en CSV.')
                for row in sheet.iter_rows():
                    values = [cell.coordinate + '=' + str(cell.value) for cell in row if cell.value is not None]
                    if values:
                        line = ' | '.join(values)
                        lines.append(line)
                        length += len(line)
                        if length > MAX_TEXT:
                            raise ValueError('Le classeur dépasse la capacité de lecture. Exporte la zone utile en CSV.')
        finally:
            workbook.close()
        return {'text': '\n'.join(lines), 'warnings': ['Valeurs et formules lues avec les références de cellules. Les formules ne sont pas recalculées ; graphiques et mise en forme non évalués.']}
    raise ValueError('Format non reconnu.')


if __name__ == '__main__':
    try:
        result = parse(Path(sys.argv[1]))
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({'error': str(error)}, ensure_ascii=False))
        sys.exit(1)
