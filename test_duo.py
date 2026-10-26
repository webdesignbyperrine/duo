import base64
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import attachments
import server
from workflows import execute, parse_review


def sample_source():
    raw = b'Le budget approuve est 1200 euros.'
    return dict(id='P-budget', name='budget.txt', raw=raw, size=len(raw), chars=len(raw), text=raw.decode(), sha256='test-hash', role='source', warnings=[])


class WorkflowTests(unittest.TestCase):
    def prepare(self):
        discussion = server.Discussion()
        discussion.active_sources = [sample_source()]
        discussion.state = {'running': True, 'messages': [{'speaker': 'Toi', 'text': 'Verifier le budget'}], 'status': '', 'error': ''}
        return discussion

    def fake(self, calls, verdict='accord', points=None, references=None):
        def call(speaker, prompt, schema=None):
            calls.append((speaker, prompt, schema))
            text = 'ANALYSE-UNIQUE-' + str(len(calls))
            if schema and 'verdict' in schema['properties']:
                return json.dumps(dict(verdict=verdict, analysis=text, open_points=points or [], source_refs=['P-budget'] if references is None else references, files=[]))
            if schema:
                return json.dumps(dict(analysis=text, files=[]))
            return text
        return call

    def test_separate_opinions_are_blind_and_share_same_sources(self):
        d = self.prepare()
        calls = []
        d.call = self.fake(calls)
        execute(d, 4, 'independent', 'Codex', d.active_sources)
        self.assertEqual(len(calls), 2)
        self.assertNotIn('ANALYSE-UNIQUE-1', calls[1][1])
        for speaker, prompt, schema in calls:
            self.assertIn('1200 euros', prompt)
        self.assertTrue(all(not m['received_artifacts'] for m in d.state['messages'][1:]))

    def test_review_passes_files_back_without_changing_original(self):
        d = self.prepare()
        calls = []
        d.call = self.fake(calls, verdict='desaccord', points=['Total incorrect'])
        execute(d, 1, 'cross_review', 'Claude Code', d.active_sources)
        self.assertEqual([c[0] for c in calls], ['Claude Code', 'Codex', 'Claude Code'])
        self.assertIn('ANALYSE-UNIQUE-1', calls[1][1])
        self.assertIn('ANALYSE-UNIQUE-2', calls[2][1])
        self.assertEqual(d.active_sources[0]['raw'], sample_source()['raw'])
        self.assertIn('Total incorrect', d.state['messages'][-1]['open_points'])
        self.assertEqual(len(d.state['messages'][-1]['received_artifacts']), 2)

    def test_convergence_round_uses_same_previous_files(self):
        d = self.prepare()
        calls = []
        d.call = self.fake(calls)
        execute(d, 4, 'convergence', 'Codex', d.active_sources)
        self.assertEqual(len(calls), 4)  # arrêt anticipé
        self.assertTrue(d.state['converged'])
        self.assertNotIn('ANALYSE-UNIQUE-1', calls[1][1])
        self.assertIn('ANALYSE-UNIQUE-1', calls[3][1])
        self.assertIn('ANALYSE-UNIQUE-2', calls[3][1])
        self.assertNotIn('ANALYSE-UNIQUE-3', calls[3][1])

    def test_reservations_and_missing_sources_prevent_convergence(self):
        for options in ({'points': ['Facture manquante']}, {'references': []}, {'verdict': 'desaccord'}):
            d = self.prepare()
            calls = []
            d.call = self.fake(calls, **options)
            execute(d, 2, 'convergence', 'Codex', d.active_sources)
            self.assertEqual(len(calls), 6)
            self.assertFalse(d.state['converged'])

    def test_stop_prevents_another_model_call(self):
        d = self.prepare()
        d.stop()
        with patch.object(d, 'call') as call:
            with self.assertRaises(InterruptedError):
                execute(d, 1, 'independent', 'Codex', d.active_sources)
            call.assert_not_called()

    def test_unknown_source_reference_cannot_declare_agreement(self):
        with self.assertRaises(RuntimeError):
            parse_review(json.dumps(dict(verdict='accord', analysis='OK', open_points=[], source_refs=['invente'])), [sample_source()])

    def test_dossier_download_contains_originals_extractions_and_analyses(self):
        d = self.prepare()
        d.call = self.fake([])
        execute(d, 1, 'independent', 'Codex', d.active_sources)
        with zipfile.ZipFile(io.BytesIO(d.bundle())) as archive:
            self.assertEqual(archive.read('pieces-originales/P-budget-budget.txt'), sample_source()['raw'])
            manifest = json.loads(archive.read('dossier-et-echanges.json'))
            self.assertEqual(manifest['sources'][0]['sha256'], 'test-hash')
            files = [n for n in archive.namelist() if n.startswith('analyses/')]
            self.assertEqual(len(files), 2)
            self.assertIn('ANALYSE-UNIQUE-1', archive.read(files[0]).decode())

    def test_upload_validation_and_no_silent_truncation(self):
        with self.assertRaises(ValueError):
            attachments.decode_upload({'name': 'piece.exe', 'content': 'AAAA'})
        with self.assertRaises(ValueError):
            attachments.decode_upload({'name': 'piece.txt', 'content': '!invalid!'})
        d = server.Discussion()
        with patch.object(server, 'extract', return_value={'text': 'x' * (attachments.MAX_TEXT + 1), 'warnings': [], 'sha256': 'x'}):
            with self.assertRaises(ValueError):
                d.upload({'name': 'piece.txt', 'content': base64.b64encode(b'abc').decode()})
        self.assertEqual(d.list_sources(), [])

    def test_401_error_is_specific(self):
        class Process:
            def __init__(self, command, **kwargs):
                self.stdin = io.StringIO()
                self.returncode = 1
                kwargs['stdout'].write(json.dumps({'is_error': True, 'api_error_status': 401}))
            def poll(self):
                return self.returncode
        with patch.object(server.subprocess, 'Popen', Process):
            with self.assertRaisesRegex(RuntimeError, 'connexion de Claude Code a expiré'):
                server.Discussion().call('Claude Code', 'Bonjour')


class ExtractionTests(unittest.TestCase):
    def test_real_file_formats(self):
        from docx import Document
        from openpyxl import Workbook
        from reportlab.pdfgen.canvas import Canvas
        from PIL import Image, ImageDraw, ImageFont
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'note.txt').write_text('Budget approuvé : 1200 euros.', encoding='utf-8')
            doc = Document()
            doc.add_paragraph('Le budget est de 1200 euros.')
            doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Facture : 950 euros'
            doc.save(root / 'note.docx')
            book = Workbook()
            book.active['A1'] = 'Budget'
            book.active['B1'] = 1200
            book.active['B2'] = '=B1-950'
            book.save(root / 'budget.xlsx')
            pdf = Canvas(str(root / 'note.pdf'))
            pdf.drawString(60, 780, 'Budget approuve : 1200 euros.')
            pdf.save()
            image = Image.new('RGB', (1500, 300), 'white')
            font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 60)
            ImageDraw.Draw(image).text((60, 80), 'Budget approuve : 1200 euros.', fill='black', font=font)
            image.save(root / 'scan.png')
            for name in ('note.txt', 'note.docx', 'budget.xlsx', 'note.pdf', 'scan.png'):
                with self.subTest(name=name):
                    result = attachments.extract(name, (root / name).read_bytes())
                    self.assertIn('1200', result['text'])
            text = attachments.extract('note.docx', (root / 'note.docx').read_bytes())['text']
            self.assertIn('950', text)
            text = attachments.extract('budget.xlsx', (root / 'budget.xlsx').read_bytes())['text']
            self.assertIn('B2==B1-950', text)
            # PDF scanné : aucune couche texte, reconnaissance automatique attendue.
            pdf = Canvas(str(root / 'scan.pdf'))
            pdf.drawImage(str(root / 'scan.png'), 20, 350, width=550, height=110)
            pdf.save()
            result = attachments.extract('scan.pdf', (root / 'scan.pdf').read_bytes())
            self.assertIn('1200', result['text'])
            self.assertTrue(any('reconnu sur image' in v for v in result['warnings']))


if __name__ == '__main__':
    unittest.main()
