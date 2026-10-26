import base64
import io
import json
from pathlib import Path
import tempfile
import threading
import unittest
import zipfile

from file_factory import build, build_files
from workflows import perform_step, execute_followup
import server


class FileTests(unittest.TestCase):
    def test_actual_binary_formats_and_recalculation(self):
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        from pypdf import PdfReader
        from PIL import Image
        document = '# Décision sur le budget\n\nLe solde disponible est de 250 euros.\n\n| Poste | Montant |\n| --- | --- |\n| Budget | 1200 |\n| Dépense | 950 |\n\n- Vérifier la facture avant règlement.'
        specs = [
            dict(name='analyse.docx', format='docx', content=document),
            dict(name='analyse.pdf', format='pdf', content=document),
            dict(name='budget.xlsx', format='xlsx', content=json.dumps({'sheets':[{'name':'Budget','rows':[['Poste','Montant'],['Budget',1200],['Dépense',950],['Solde',None]],'formulas':[{'cell':'B4','formula':'=B2-B3'}]}]})),
            dict(name='decision.pptx', format='pptx', content=json.dumps({'slides':[{'title':'Décision sur le budget','body':'Le budget approuvé est de 1200 euros.','bullets':['Dépense prévue : 950 euros','Solde disponible : 250 euros']}]})),
            dict(name='schema.png', format='png', content=json.dumps({'width':1200,'height':500,'elements':[{'type':'text','x':70,'y':70,'text':'Budget : 1200 euros','size':44}]})),
        ]
        with tempfile.TemporaryDirectory() as folder:
            for spec in specs:
                with self.subTest(format=spec['format']):
                    result = build(spec)
                    path = Path(folder)/spec['name'];path.write_bytes(result['raw'])
                    self.assertGreater(result['size'], 100)
                    self.assertIn('1200', result['review_text'])
                    if spec['format']=='docx':self.assertTrue(Document(path).tables)
                    if spec['format']=='pdf':self.assertIn('250 euros', '\n'.join(p.extract_text() for p in PdfReader(path).pages))
                    if spec['format']=='xlsx':
                        values=load_workbook(path,data_only=True);self.assertEqual(values['Budget']['B4'].value,250);values.close()
                        self.assertIn('Budget!B4 = 250',result['review_text'])
                        self.assertFalse(any('formules ne sont pas recalculées' in warning for warning in result['warnings']))
                    if spec['format']=='pptx':self.assertEqual(len(Presentation(path).slides),1)
                    if spec['format']=='png':
                        with Image.open(path) as image:self.assertEqual(image.size,(1200,500))

    def test_archive_and_filename_do_not_masquerade(self):
        result=build({'name':'../../rapport.exe','format':'pdf','content':'# Rapport\n\nConclusion.'})
        self.assertEqual(result['name'],'rapport.pdf');self.assertTrue(result['raw'].startswith(b'%PDF'))
        result=build({'name':'dossier.zip','format':'zip','content':json.dumps({'files':[{'name':'notes.md','format':'md','content':'Budget confirmé : 1200 euros.'},{'name':'donnees.json','format':'json','content':'{"budget":1200}'}]})})
        with zipfile.ZipFile(io.BytesIO(result['raw'])) as archive:
            self.assertEqual(set(archive.namelist()),{'notes.md','donnees.json'})
            self.assertEqual(json.loads(archive.read('donnees.json'))['budget'],1200)

    def test_failed_creation_never_gets_download_link(self):
        files,errors=build_files([{'name':'faux.dwg','format':'dwg','content':'pas un dessin CAD'}])
        self.assertFalse(files);self.assertTrue(errors)
        files,errors=build_files([{'name':'table.xlsx','format':'xlsx','content':'{"sheets":[]}'}])
        self.assertFalse(files);self.assertTrue(errors)

    def test_external_spreadsheet_formula_is_rejected(self):
        spec={'name':'table.xlsx','format':'xlsx','content':json.dumps({'sheets':[{'rows':[['Donnée'],['=WEBSERVICE("https://example.com")']]}]})}
        with self.assertRaisesRegex(ValueError,'accès externe'):build(spec)

    def test_generated_bytes_and_extracted_content_are_passed_to_next_step(self):
        d=server.Discussion();d.state={'running':True,'messages':[{'speaker':'Toi','text':'Crée un fichier de budget'}],'status':'','error':''}
        first=d.append_result('Codex','Création','Le fichier est préparé.',[],file_specs=[{'name':'budget.txt','format':'txt','content':'VALEUR_DANS_FICHIER=1200'}])
        observed=[]
        def call(speaker,prompt,schema):
            observed.append(prompt)
            return json.dumps({'analysis':'Le contenu reçu indique 1200.','verdict':'accord','open_points':[],'source_refs':[],'files':[]})
        d.call=call
        second=perform_step(d,'Claude Code','Relecture','Vérifie le fichier',[],[first],'Relis les fichiers.',True)
        self.assertIn('VALEUR_DANS_FICHIER=1200',observed[0])
        self.assertIn(first['outputs'][0]['sha256'],observed[0])
        self.assertEqual(second['received_files'],[first['outputs'][0]['id']])
        with zipfile.ZipFile(io.BytesIO(d.bundle())) as archive:
            name=next(n for n in archive.namelist() if n.startswith('livrables/'))
            self.assertEqual(archive.read(name),b'VALEUR_DANS_FICHIER=1200')

    def test_followup_preserves_previous_messages_and_files(self):
        d=server.Discussion();d.state={'running':True,'messages':[{'speaker':'Toi','text':'Budget de 1200 euros.'}],'status':'','error':'','first':'Codex'}
        previous=d.append_result('Codex','Analyse','Solde confirmé.',[],file_specs=[{'name':'resultat.txt','format':'txt','content':'solde=250'}])
        seen=[]
        def call(speaker,prompt,schema):
            seen.append((speaker,prompt))
            value={'analysis':'Fichier vérifié.','files':[]}
            if 'verdict' in schema['properties']:value.update(verdict='accord',open_points=[],source_refs=[])
            return json.dumps(value)
        d.call=call
        execute_followup(d,'Prépare un fichier et fais-le vérifier.','both')
        self.assertEqual([x[0] for x in seen],['Codex','Claude Code'])
        self.assertIn('solde=250',seen[1][1]);self.assertIn(previous['outputs'][0]['id'],d.generated)
        self.assertEqual(len(d.state['messages']),4)

    def test_stop_cancels_file_creation(self):
        event=threading.Event();event.set()
        with self.assertRaises(InterruptedError):build_files([{'name':'test.txt','format':'txt','content':'test'}],event)


if __name__=='__main__':unittest.main()
